"""Fase 8 aislada: --pruebas, --http y --abrir. No usa la SQLite del negocio."""
import argparse
import hashlib
import json
import os
from pathlib import Path
import secrets
import subprocess
import sys
import tempfile
from verificar_catalogo_fase3 import preparar_entorno

RAIZ = Path(__file__).resolve().parent.parent
PROTEGIDOS = ('backend/datos/db.sqlite3', 'backend/.env',
              'backend/REPORTE_MIGRACIONES_FASE1.md', 'backend/SOLUCION_ERROR_TERMINAL.md')


def huellas():
    return {p: hashlib.sha256((RAIZ / p).read_bytes()).hexdigest() if (RAIZ / p).exists() else None for p in PROTEGIDOS}


def recorrido_http(temporal, abrir):
    import threading
    from collections import Counter
    from decimal import Decimal
    from html.parser import HTMLParser
    from http.cookiejar import CookieJar
    from io import BytesIO
    from urllib.error import HTTPError
    from urllib.parse import urlencode
    from urllib.request import HTTPCookieProcessor, Request, build_opener
    from wsgiref.simple_server import make_server, WSGIRequestHandler
    import django
    django.setup()
    from django.conf import settings
    from django.core.management import call_command
    from django.core.wsgi import get_wsgi_application
    from django.db import connections
    from openpyxl import load_workbook
    from apps.core.frontend import RECURSOS_FRONTEND
    from apps.reportes.services import REPORTES, consultar_reporte
    from apps.reportes.tests.datos import crear_datos
    from verificar_interfaz_fase0 import TokenCSRF

    if Path(settings.DATABASES['default']['NAME']).resolve() != (temporal / 'aislada.sqlite3').resolve():
        raise RuntimeError('Se rechazó una base distinta de la temporal.')
    call_command('migrate', verbosity=0, interactive=False)
    password = secrets.token_urlsafe(24)
    datos = crear_datos(password)
    app = get_wsgi_application()
    peticiones = []
    def sitio(environ, start_response):
        def responder(status, headers, exc_info=None):
            peticiones.append(dict(ruta=environ['PATH_INFO'], estado=int(status[:3])))
            return start_response(status, headers, exc_info)
        return app(environ, responder)
    class Silencioso(WSGIRequestHandler):
        def log_message(self, *args):
            pass
    servidor = make_server('127.0.0.1', 0, sitio, handler_class=Silencioso)
    def atender():
        try:
            servidor.serve_forever()
        finally:
            connections.close_all()
    hilo = threading.Thread(target=atender, daemon=True)
    hilo.start()
    origen = f'http://127.0.0.1:{servidor.server_port}'
    informe = dict(aprobado=False, visual='Pendiente de inspección manual', peticiones=peticiones, exportaciones={})
    def pedir(cliente, ruta, post=None):
        solicitud = Request(origen + ruta, data=urlencode(post).encode() if post is not None else None)
        try:
            respuesta = cliente.open(solicitud, timeout=30)
        except HTTPError as error:
            respuesta = error
        with respuesta:
            return respuesta.status, respuesta.read(), dict(respuesta.headers)
    def sesion(nombre):
        cliente = build_opener(HTTPCookieProcessor(CookieJar()))
        estado, html, _ = pedir(cliente, '/login/')
        assert estado == 200
        token = TokenCSRF(html.decode())
        assert pedir(cliente, '/login/', dict(username=nombre, password=password, csrfmiddlewaretoken=token.token))[0] == 200
        return cliente
    class Recursos(HTMLParser):
        def __init__(self):
            super().__init__()
            self.rutas = set()
        def handle_starttag(self, tag, attrs):
            if tag not in ('script', 'link', 'img', 'iframe'):
                return
            atributos = dict(attrs)
            ruta = atributos.get('src') or atributos.get('href')
            if ruta:
                if not ruta.startswith('/static/'):
                    raise AssertionError(f'Recurso no local: {ruta}')
                self.rutas.add(ruta)
    try:
        dueno, vendedor = sesion('dueno_fase8'), sesion('vendedor_fase8')
        assert pedir(dueno, '/reportes/')[0] == 200
        recursos = Recursos()
        for nombre in REPORTES:
            filtros = {} if nombre == 'inventario_valorizado' else {'dias': 30} if nombre == 'vencimientos' else {'desde': datos['hoy'].isoformat(), 'hasta': datos['hoy'].isoformat()}
            if nombre == 'ventas':
                filtros['agrupar'] = 'dia'
            consulta = urlencode(filtros)
            estado, html, _ = pedir(dueno, f'/reportes/{nombre}/?{consulta}')
            assert estado == 200, (nombre, estado)
            recursos.feed(html.decode())
            (temporal / f'{nombre}.html').write_bytes(html)
            esperado = consultar_reporte(datos['dueno'], nombre, filtros)
            for formato in ('xlsx', 'pdf'):
                url = f'/reportes/{nombre}/exportar/?' + urlencode({**filtros, 'formato': formato})
                estado, archivo, cabeceras = pedir(dueno, url)
                assert estado == 200 and 'attachment' in cabeceras['Content-Disposition']
                if formato == 'xlsx':
                    hoja = load_workbook(BytesIO(archivo)).active
                    assert hoja.max_row == len(esperado['filas']) + 2
                    for indice, fila in enumerate(esperado['filas'], 2):
                        for col, (clave, _, _) in enumerate(esperado['columnas'], 1):
                            valor = hoja.cell(indice, col).value
                            assert (Decimal(str(valor)) if isinstance(fila[clave], Decimal) else valor) == fila[clave]
                else:
                    assert archivo.startswith(b'%PDF-') and b'%%EOF' in archivo
                informe['exportaciones'][nombre + '.' + formato] = dict(bytes=len(archivo), sha256=hashlib.sha256(archivo).hexdigest())
                assert pedir(vendedor, url)[0] == (200 if nombre == 'vencimientos' else 403)
            assert pedir(vendedor, f'/reportes/{nombre}/?{consulta}')[0] == (200 if nombre == 'vencimientos' else 403)
        assert pedir(dueno, '/reportes/ventas/?desde=no')[0] == 422
        assert pedir(dueno, '/reportes/ventas/exportar/?formato=exe')[0] == 422
        for ruta in recursos.rutas | {'/static/' + r for r in RECURSOS_FRONTEND}:
            assert pedir(dueno, ruta)[0] == 200
        assert not any(p['estado'] in (404, 500) or p['estado'] > 500 for p in peticiones)
        informe['aprobado'] = True
        informe['estados'] = dict(Counter(p['estado'] for p in peticiones))
        print(f'HTTP aprobado: {len(peticiones)} solicitudes; ocho reportes y dieciséis exportaciones.', flush=True)
        if abrir:
            import webbrowser
            print(f'Revisión visual: {origen}/reportes/', flush=True)
            print(f'Usuarios temporales: dueno_fase8 / vendedor_fase8. Contraseña: {password}', flush=True)
            print(f'Rango con datos ficticios: {datos["hoy"].isoformat()} a {datos["hoy"].isoformat()}. Revisar 1280x800 y 390x844. Ctrl+C finaliza.', flush=True)
            webbrowser.open(origen + '/reportes/')
            try:
                threading.Event().wait()
            except KeyboardInterrupt:
                pass
    finally:
        servidor.shutdown()
        servidor.server_close()
        hilo.join(timeout=5)
        connections.close_all()
        (temporal / 'http.json').write_text(json.dumps(informe, ensure_ascii=False, indent=2), encoding='utf-8')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for opcion in ('pruebas', 'http', 'abrir'):
        parser.add_argument('--' + opcion, action='store_true')
    args = parser.parse_args()
    if not any(vars(args).values()):
        parser.error('Elige --pruebas, --http o --abrir.')
    if Path(sys.executable).resolve() != (RAIZ / 'venv/Scripts/python.exe').resolve():
        parser.error('Utiliza exclusivamente venv/Scripts/python.exe.')
    antes = huellas()
    temporal = Path(tempfile.mkdtemp(prefix='Inventarios_verificacion_fase8_'))
    env = preparar_entorno(temporal)
    env.update(NOMBRE_NEGOCIO='Reportes de demostración temporal', TIME_ZONE='America/La_Paz')
    resultados = {}
    print(f'Evidencia y SQLite aisladas: {temporal}', flush=True)
    try:
        if args.pruebas:
            comandos = {
                'check': ['manage.py', 'check'],
                'makemigrations': ['manage.py', 'makemigrations', '--check', '--dry-run'],
                'test': ['manage.py', 'test', '--noinput'],
                'frontend': [str(RAIZ / 'scripts/preparar_frontend.py'), '--verificar'],
            }
            for nombre, comando in comandos.items():
                with (temporal / f'{nombre}.txt').open('w', encoding='utf-8') as salida:
                    proceso = subprocess.run([sys.executable, *comando], cwd=RAIZ / 'backend', env=env, stdout=salida, stderr=subprocess.STDOUT)
                resultados[nombre] = proceso.returncode
                if nombre == 'test' and 'Ran 0 tests' in (temporal / 'test.txt').read_text(encoding='utf-8'):
                    resultados[nombre] = 1
                print(f'{nombre}: código {resultados[nombre]}', flush=True)
        if args.http or args.abrir:
            os.environ.update(env)
            sys.path[:0] = [str(temporal), str(RAIZ / 'backend')]
            recorrido_http(temporal, args.abrir)
            resultados['http'] = 0
    finally:
        resultados['archivos_protegidos_intactos'] = antes == huellas()
        (temporal / 'resultados.json').write_text(json.dumps(resultados, indent=2), encoding='utf-8')
        print('Base real, .env y reportes intactos:', resultados['archivos_protegidos_intactos'], flush=True)
    return int(not resultados.pop('archivos_protegidos_intactos') or any(resultados.values()))


if __name__ == '__main__':
    sys.exit(main())
