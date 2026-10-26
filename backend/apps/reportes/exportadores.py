"""Documentos en memoria; solo formatean resultados autorizados del servicio."""
from io import BytesIO
from xml.sax.saxutils import escape
from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
from .services import autorizar_reporte


def texto_celda(valor, tipo):
    if valor is None:
        return '—'
    return format(valor, '.2f') if tipo == 'moneda' else str(valor)


def matriz(reporte):
    """Misma proyección ordenada para las tres representaciones."""
    return [[fila[k] for k, _, _ in reporte['columnas']] for fila in reporte['filas']]


def exportar_excel(reporte, *, usuario):
    """Genera Excel sin archivos persistentes.

    Pasos: autorizar; escribir datos y totales; serializar en memoria.
    Lanza: PermissionDenied si el actor perdió permisos.
    """
    # Paso 1: las descargas también verifican permisos al invocar el exportador.
    autorizar_reporte(usuario, reporte['nombre'])
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Reporte'
    encabezados = [c[1] for c in reporte['columnas']]
    # Paso 2: forzar cadenas como texto, incluso si empiezan con =, +, - o @.
    def escribir(valores):
        hoja.append(valores)
        for celda in hoja[hoja.max_row]:
            if isinstance(celda.value, str):
                celda.data_type = 's'
    escribir(encabezados)
    for fila in matriz(reporte):
        escribir(fila)
    escribir(['Totales'] + [reporte['totales'].get(k) for k, _, _ in reporte['columnas'][1:]])
    for fila in (hoja[1], hoja[hoja.max_row]):
        for celda in fila:
            celda.font = Font(bold=True)
    for indice, (_, _, tipo) in enumerate(reporte['columnas'], 1):
        hoja.column_dimensions[get_column_letter(indice)].width = min(60, max(14, max(len(str(c.value or '')) for c in hoja[get_column_letter(indice)]) + 2))
        if tipo == 'moneda':
            for fila in hoja.iter_rows(min_row=2, min_col=indice, max_col=indice):
                fila[0].number_format = '#,##0.00'
    hoja.freeze_panes = 'A2'
    meta = libro.create_sheet('Información')
    for clave, valor in [('Negocio', settings.NOMBRE_NEGOCIO), ('Reporte', reporte['titulo']), ('Moneda', settings.MONEDA_SIMBOLO), *reporte['filtros'].items(), *[('Nota', n) for n in reporte['notas']]]:
        meta.append([clave, valor])
        for c in meta[meta.max_row]:
            c.data_type = 's'
    meta.column_dimensions['A'].width = 22
    meta.column_dimensions['B'].width = 90
    # Paso 3: nunca escribir en DATA_DIR ni guardar el reporte en el servidor.
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def exportar_pdf(reporte, *, usuario):
    """Genera PDF con tablas largas y encabezado por página.

    Pasos: autorizar; construir tabla; paginar en memoria.
    Lanza: PermissionDenied si el actor perdió permisos.
    """
    # Paso 1: no inferir autorización del enlace o de la pantalla anterior.
    autorizar_reporte(usuario, reporte['nombre'])
    salida = BytesIO()
    ancho, alto = landscape(A4) if len(reporte['columnas']) > 5 else A4
    documento = SimpleDocTemplate(salida, pagesize=(ancho, alto), leftMargin=28, rightMargin=28, topMargin=54, bottomMargin=36, title=reporte['titulo'])
    estilos = getSampleStyleSheet()
    estilos['BodyText'].fontSize = 8
    estilos['BodyText'].leading = 11
    parrafo = lambda texto: Paragraph(escape(str(texto)), estilos['BodyText'])
    historia = [Paragraph(escape(reporte['titulo']), estilos['Heading1']), parrafo(' · '.join(f'{k}: {v}' for k, v in reporte['filtros'].items()) or 'Todos los registros aplicables'), parrafo('Moneda: ' + settings.MONEDA_SIMBOLO), Spacer(1, 10)]
    # Paso 2: texto escapado y encabezados repetidos; no dividir una fila común.
    tabla = [[parrafo(c[1]) for c in reporte['columnas']]]
    for fila in matriz(reporte):
        tabla.append([parrafo(texto_celda(v, c[2])) for v, c in zip(fila, reporte['columnas'])])
    tabla.append([parrafo('Totales')] + [parrafo(texto_celda(reporte['totales'].get(k), t)) for k, _, t in reporte['columnas'][1:]])
    componente = LongTable(tabla, colWidths=[(ancho - 56) / len(reporte['columnas'])] * len(reporte['columnas']), repeatRows=1, splitByRow=1)
    componente.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e9ecef')), ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#e9ecef')), ('GRID', (0, 0), (-1, -1), .3, colors.grey), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 5), ('RIGHTPADDING', (0, 0), (-1, -1), 5)]))
    historia.append(componente)
    if not reporte['filas']:
        historia.append(parrafo('No hay resultados para los filtros seleccionados.'))
    for nota in reporte['notas']:
        historia.extend([Spacer(1, 8), parrafo(nota)])
    # Paso 3: cada página conserva negocio y numeración.
    def pagina(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 9)
        canvas.drawString(28, alto - 30, settings.NOMBRE_NEGOCIO)
        canvas.drawRightString(ancho - 28, 20, f'Página {doc.page}')
        canvas.restoreState()
    documento.build(historia, onFirstPage=pagina, onLaterPages=pagina)
    return salida.getvalue()
