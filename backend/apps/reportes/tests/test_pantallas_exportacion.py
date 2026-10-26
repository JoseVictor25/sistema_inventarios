"""Permisos en tres capas y contenido real de los documentos descargados."""
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from django.contrib.auth.models import AnonymousUser
from django.core.exceptions import PermissionDenied
from django.test import TestCase, Client
from django.urls import reverse
from openpyxl import load_workbook
from reportlab.platypus import LongTable
from apps.reportes import services, exportadores
from apps.usuarios.models import Usuario, Rol, Permiso
from .datos import crear_datos


class ReportesPantallasTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.d = crear_datos()

    def filtros(self, nombre):
        if nombre == 'inventario_valorizado':
            return {}
        if nombre == 'vencimientos':
            return {'dias': 30}
        datos = {'desde': self.d['hoy'].isoformat(), 'hasta': self.d['hoy'].isoformat()}
        if nombre == 'ventas':
            datos['agrupar'] = 'dia'
        return datos

    def test_ocho_pantallas_y_menu(self):
        self.client.force_login(self.d['dueno'])
        self.assertContains(self.client.get('/reportes/'), 'Reportes y exportación')
        for nombre in services.REPORTES:
            with self.subTest(nombre=nombre):
                respuesta = self.client.get(reverse('reportes:' + nombre), self.filtros(nombre))
                self.assertEqual(respuesta.status_code, 200)
                self.assertContains(respuesta, 'Descargar Excel')
                self.assertContains(respuesta, 'Descargar PDF')
                self.assertTemplateUsed(respuesta, f'reportes/{nombre}.html')
                self.assertEqual(respuesta.context['filas'], services.consultar_reporte(self.d['dueno'], nombre, self.filtros(nombre))['filas'])

    def test_vendedor_solo_vencimientos_y_exportacion(self):
        self.client.force_login(self.d['vendedor'])
        menu = self.client.get('/reportes/')
        self.assertContains(menu, 'Vencimientos próximos')
        self.assertNotContains(menu, 'Ganancia bruta')
        for nombre in services.REPORTES:
            for formato in (None, 'xlsx', 'pdf'):
                with self.subTest(nombre=nombre, formato=formato):
                    datos = self.filtros(nombre)
                    if formato:
                        datos['formato'] = formato
                    url = reverse('reportes:' + nombre + ('_exportar' if formato else ''))
                    respuesta = self.client.get(url, datos, HTTP_HX_REQUEST='true', HTTP_ACCEPT='application/json')
                    self.assertEqual(respuesta.status_code, 200 if nombre == 'vencimientos' else 403)

    def test_anonimo_inactivo_sin_rol_superusuario(self):
        sinrol = Usuario.objects.create_user('sinrol')
        inactivo = Usuario.objects.create_user('inactivo', is_active=False, rol=self.d['dueno'].rol)
        superusuario = Usuario.objects.create_superuser('super', password='Temporal-super-2026!')
        for usuario in [AnonymousUser(), sinrol, inactivo]:
            for nombre in services.REPORTES:
                with self.subTest(usuario=str(usuario), reporte=nombre), self.assertRaises(PermissionDenied):
                    services.consultar_reporte(usuario, nombre, self.filtros(nombre))
        for nombre in services.REPORTES:
            self.assertEqual(self.client.get(reverse('reportes:' + nombre)).status_code, 302)
            services.consultar_reporte(superusuario, nombre, self.filtros(nombre))
        self.client.force_login(inactivo)
        self.assertEqual(self.client.get('/reportes/').status_code, 302)
        self.client.force_login(sinrol)
        self.assertEqual(self.client.get('/reportes/').status_code, 403)

    def test_ver_reportes_no_concede_costos(self):
        rol = Rol.objects.create(nombre='Prueba solo reportes')
        rol.permisos.add(Permiso.objects.get(codigo='ver_reportes'))
        usuario = Usuario.objects.create_user('consultor', rol=rol)
        self.client.force_login(usuario)
        for nombre in ('compras', 'inventario_valorizado', 'ganancias'):
            with self.assertRaises(PermissionDenied):
                services.consultar_reporte(usuario, nombre, self.filtros(nombre))
            for sufijo in ('', '_exportar'):
                self.assertEqual(self.client.get(reverse('reportes:' + nombre + sufijo), {**self.filtros(nombre), 'formato': 'xlsx'}).status_code, 403)
        rol.permisos.clear()
        rol.permisos.add(Permiso.objects.get(codigo='ver_costos'))
        with self.assertRaises(PermissionDenied):
            services.consultar_reporte(usuario, 'ganancias', self.filtros('ganancias'))

    def test_exportadores_reautorizan(self):
        resultado = services.consultar_reporte(self.d['dueno'], 'ganancias', self.filtros('ganancias'))
        for exportador in (exportadores.exportar_excel, exportadores.exportar_pdf):
            with self.assertRaises(PermissionDenied):
                exportador(resultado, usuario=self.d['vendedor'])
            self.d['dueno'].is_active = False
            with self.assertRaises(PermissionDenied):
                exportador(resultado, usuario=self.d['dueno'])

    def test_filtros_y_formato_invalidos_http(self):
        self.client.force_login(self.d['dueno'])
        self.assertEqual(self.client.get('/reportes/ventas/').status_code, 200)
        for url in ['/reportes/ventas/?desde=no', '/reportes/vencimientos/?dias=-1', '/reportes/ventas/?desde=2026-01-01&desde=2026-01-02', '/reportes/ventas/exportar/?formato=exe', '/reportes/ventas/exportar/?formato=pdf&formato=xlsx', '/reportes/inventario_valorizado/?extra=1']:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 422)

    def test_metodos_http_solo_lectura_csrf(self):
        self.client.force_login(self.d['dueno'])
        for nombre in services.REPORTES:
            self.assertEqual(self.client.post(reverse('reportes:' + nombre)).status_code, 405)
        estricto = Client(enforce_csrf_checks=True)
        estricto.force_login(self.d['dueno'])
        self.assertEqual(estricto.post('/reportes/ventas/').status_code, 403)

    def test_excel_ocho_reportes_coinciden_con_servicio(self):
        for nombre in services.REPORTES:
            with self.subTest(nombre=nombre):
                resultado = services.consultar_reporte(self.d['dueno'], nombre, self.filtros(nombre))
                contenido = exportadores.exportar_excel(resultado, usuario=self.d['dueno'])
                hoja = load_workbook(BytesIO(contenido)).active
                self.assertEqual([c.value for c in hoja[1]], [c[1] for c in resultado['columnas']])
                self.assertTrue(all(c.font.bold for c in hoja[1]))
                for indice, fila in enumerate(exportadores.matriz(resultado), 2):
                    for columna, valor in enumerate(fila, 1):
                        obtenido = hoja.cell(indice, columna).value
                        self.assertEqual(Decimal(str(obtenido)) if isinstance(valor, Decimal) else obtenido, valor)
                        if resultado['columnas'][columna - 1][2] == 'moneda':
                            self.assertEqual(hoja.cell(indice, columna).number_format, '#,##0.00')
                self.assertEqual(hoja.cell(hoja.max_row, 1).value, 'Totales')
                for indice, (clave, _, _) in enumerate(resultado['columnas'][1:], 2):
                    valor = resultado['totales'].get(clave)
                    obtenido = hoja.cell(hoja.max_row, indice).value
                    self.assertEqual(Decimal(str(obtenido)) if isinstance(valor, Decimal) else obtenido, valor)

    def test_pdf_ocho_reportes_tablas_y_totales(self):
        for nombre in services.REPORTES:
            with self.subTest(nombre=nombre):
                resultado = services.consultar_reporte(self.d['dueno'], nombre, self.filtros(nombre))
                with patch('apps.reportes.exportadores.LongTable', wraps=LongTable) as tabla:
                    contenido = exportadores.exportar_pdf(resultado, usuario=self.d['dueno'])
                self.assertTrue(contenido.startswith(b'%PDF-'))
                self.assertIn(b'%%EOF', contenido)
                datos = tabla.call_args.args[0]
                for fila_pdf, fila in zip(datos[1:-1], exportadores.matriz(resultado)):
                    self.assertEqual([p.getPlainText() for p in fila_pdf], [exportadores.texto_celda(v, c[2]) for v, c in zip(fila, resultado['columnas'])])
                self.assertEqual(datos[-1][0].getPlainText(), 'Totales')
                self.assertEqual(tabla.call_args.kwargs['repeatRows'], 1)

    def test_documentos_vacios(self):
        resultado = services.consultar_reporte(self.d['dueno'], 'ventas', {'desde': '2000-01-01', 'hasta': '2000-01-01', 'agrupar': 'dia'})
        self.assertEqual(load_workbook(BytesIO(exportadores.exportar_excel(resultado, usuario=self.d['dueno']))).active.max_row, 2)
        self.assertTrue(exportadores.exportar_pdf(resultado, usuario=self.d['dueno']).startswith(b'%PDF-'))

    def test_excel_inyeccion_formula_y_pdf_texto_escapado(self):
        resultado = services.consultar_reporte(self.d['dueno'], 'productos_mas_vendidos', self.filtros('productos_mas_vendidos'))
        resultado['filas'][0]['producto'] = '=HYPERLINK("x") <b>Texto</b> & dato'
        hoja = load_workbook(BytesIO(exportadores.exportar_excel(resultado, usuario=self.d['dueno']))).active
        self.assertEqual(hoja['A2'].data_type, 's')
        self.assertEqual(hoja['A2'].value, resultado['filas'][0]['producto'])
        self.assertTrue(exportadores.exportar_pdf(resultado, usuario=self.d['dueno']).startswith(b'%PDF-'))

    def test_pdf_largo_y_excel_sin_recorte(self):
        resultado = services.consultar_reporte(self.d['dueno'], 'productos_mas_vendidos', self.filtros('productos_mas_vendidos'))
        resultado['filas'] = [dict(producto='Producto largo de prueba ' * 5, unidades=1, total=Decimal('1.25')) for _ in range(300)]
        resultado['totales'] = {'unidades': 300, 'total': Decimal('375')}
        pdf = exportadores.exportar_pdf(resultado, usuario=self.d['dueno'])
        self.assertGreater(pdf.count(b'/Type /Page\n'), 1)
        self.assertEqual(load_workbook(BytesIO(exportadores.exportar_excel(resultado, usuario=self.d['dueno']))).active.max_row, 302)

    def test_descargas_http_contenido_y_nombres(self):
        self.client.force_login(self.d['dueno'])
        for nombre in services.REPORTES:
            for formato in ('xlsx', 'pdf'):
                with self.subTest(nombre=nombre, formato=formato):
                    respuesta = self.client.get(reverse('reportes:' + nombre + '_exportar'), {**self.filtros(nombre), 'formato': formato})
                    self.assertEqual(respuesta.status_code, 200)
                    self.assertIn(nombre, respuesta['Content-Disposition'])
                    self.assertIn('no-store', respuesta['Cache-Control'])
                    self.assertTrue(respuesta.content.startswith(b'PK' if formato == 'xlsx' else b'%PDF-'))

    def test_grafico_local_serializado_seguro(self):
        self.client.force_login(self.d['dueno'])
        respuesta = self.client.get('/reportes/ventas/', self.filtros('ventas'))
        self.assertContains(respuesta, 'id="datos-grafico"')
        self.assertContains(respuesta, '/static/js/reportes.js')
        self.assertContains(respuesta, '/static/vendor/chartjs/chart.umd.min.js')
        self.assertNotContains(respuesta, 'src="https://')
